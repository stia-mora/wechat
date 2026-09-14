"""Export stored PostgreSQL articles; never contact an acquisition service."""

import html
import io
import json
import zipfile

from fastapi import HTTPException
from fastapi.responses import Response

from .db import db


def export_account(account_id, format):
    if format not in ("zip", "html", "json", "xlsx", "docx", "pdf", "epub"):
        raise HTTPException(422, "不支持的导出格式")
    with db() as conn:
        account = conn.execute(
            "SELECT name FROM official_accounts WHERE id=%s", (account_id,)
        ).fetchone()
        if not account:
            raise HTTPException(404, "公众号不存在")
        articles = conn.execute(
            "SELECT * FROM articles WHERE account_id=%s AND content_text<>'' ORDER BY publish_time DESC NULLS LAST,id DESC LIMIT 3000",
            (account_id,),
        ).fetchall()
    if not articles:
        raise HTTPException(404, "本站还没有可导出的正文，请先完成采集")
    title = account["name"]
    buffer = io.BytesIO()
    metadata = [
        {
            "title": a["title"],
            "author": a["author"],
            "publish_time": a["publish_time"].isoformat() if a["publish_time"] else None,
            "source_url": a["source_url"],
            "cover_url": a["cover_url"],
        }
        for a in articles
    ]
    if format == "json":
        buffer.write(json.dumps(metadata, ensure_ascii=False, indent=2).encode())
        mime = "application/json"
    elif format == "html":
        sections = "".join(
            f'<article><h2>{html.escape(a["title"])}</h2><p><a href="{html.escape(a["source_url"], quote=True)}">原文</a></p>{a["content_html"]}</article>'
            for a in articles
        )
        buffer.write(
            (
                '<!doctype html><meta charset="utf-8"><title>'
                + html.escape(title)
                + "</title><style>body{max-width:800px;margin:40px auto;padding:20px;line-height:1.8}img{max-width:100%}article{margin-bottom:60px}</style><h1>"
                + html.escape(title)
                + "</h1>"
                + sections
            ).encode()
        )
        mime = "text/html"
    elif format == "zip":
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for a in articles:
                archive.writestr(
                    f"{a['id']}.md",
                    f"# {a['title']}\n\n作者：{a['author']}\n\n原文：{a['source_url']}\n\n{a['content_text']}\n",
                )
        mime = "application/zip"
    elif format == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(list(metadata[0]))
        for row in metadata:
            values = [str(v) if v is not None else "" for v in row.values()]
            sheet.append(["'" + v if v.startswith(("=", "+", "-", "@")) else v for v in values])
        book.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading(title, 0)
        for a in articles[:500]:
            doc.add_heading(a["title"], 1)
            doc.add_paragraph(a["source_url"])
            doc.add_paragraph(a["content_text"])
        doc.save(buffer)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format == "pdf":
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle("Chinese", fontName="STSong-Light", fontSize=11, leading=18)
        flow = []
        for a in articles[:200]:
            flow.extend([Paragraph(html.escape(a["title"]), style), Spacer(1, 12)])
            for line in a["content_text"].splitlines():
                if line:
                    flow.append(Paragraph(html.escape(line), style))
            flow.append(Spacer(1, 24))
        SimpleDocTemplate(buffer).build(flow)
        mime = "application/pdf"
    else:
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier(f"wechat-source-{account_id}")
        book.set_title(title)
        book.set_language("zh")
        chapters = []
        for a in articles[:500]:
            chapter = epub.EpubHtml(title=a["title"], file_name=f"{a['id']}.xhtml", lang="zh")
            chapter.content = (
                "<h1>"
                + html.escape(a["title"])
                + "</h1><p>"
                + html.escape(a["content_text"]).replace("\n", "</p><p>")
                + "</p>"
            )
            book.add_item(chapter)
            chapters.append(chapter)
        book.toc = chapters
        book.spine = ["nav", *chapters]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(buffer, book)
        mime = "application/epub+zip"
    return Response(
        buffer.getvalue(),
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="account-{account_id}.{format}"',
            "Cache-Control": "no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )
